#!/usr/bin/env python3
"""
universal_ingestion_tool_with_feedback_v2.py (CLI-compatible shim)

Goal: keep the Streamlit app/pipeline working on Streamlit Cloud even when
full scraping/extractor logic is evolving.

This script:
- Accepts the arguments the app passes (feedback-db, min-confidence, rules, etc.)
- Opens the WEX skeleton XLSX
- Fills Part Name + Part Description using simple heuristics + optional rule files
- Writes a Review sheet with confidence + source
- Optionally persists observations in a tiny SQLite DB (feedback-db)
- Optionally applies overrides from that DB before writing results

You can replace this file later with a richer scraper. The interface stays stable.
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Optional, Tuple, List

import openpyxl


DEFAULT_MODEL_COL = "AD"   # Part Model Number in your WEX Single Part template
DEFAULT_NAME_COL = "V"     # Part Name
DEFAULT_DESC_COL = "AE"    # Part_Description


@dataclass
class RulePack:
    warranty_text_by_brand: Dict[str, str]
    series_map: Dict[str, Dict[str, Any]]  # prefix -> attributes


def _load_json(path: Optional[str]) -> Any:
    if not path:
        return None
    p = Path(path)
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None


def load_rules(brand: str, warranty_rules_path: Optional[str], series_rules_path: Optional[str]) -> RulePack:
    warranty = _load_json(warranty_rules_path)
    series = _load_json(series_rules_path)

    warranty_text_by_brand: Dict[str, str] = {}
    # Support formats: {"Bryant": "..."} OR [{"brand":"Bryant","text":"..."}]
    if isinstance(warranty, dict):
        for k, v in warranty.items():
            if isinstance(v, str):
                warranty_text_by_brand[str(k).strip().lower()] = v.strip()
    elif isinstance(warranty, list):
        for row in warranty:
            if isinstance(row, dict):
                b = str(row.get("brand", "")).strip().lower()
                t = str(row.get("text", "")).strip()
                if b and t:
                    warranty_text_by_brand[b] = t

    series_map: Dict[str, Dict[str, Any]] = {}
    # Support formats: {"987MC": {"equipment_type":"Gas Furnace", ...}, ...}
    # OR [{"prefix":"987MC","equipment_type":"Gas Furnace", ...}, ...]
    if isinstance(series, dict):
        for k, v in series.items():
            if isinstance(v, dict):
                series_map[str(k).strip().upper()] = v
    elif isinstance(series, list):
        for row in series:
            if isinstance(row, dict):
                pref = str(row.get("prefix", "")).strip().upper()
                if pref:
                    series_map[pref] = row

    return RulePack(warranty_text_by_brand=warranty_text_by_brand, series_map=series_map)


def fb_connect(path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(path)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS observations (
            model TEXT PRIMARY KEY,
            name TEXT,
            description TEXT,
            confidence REAL,
            source TEXT,
            updated_utc TEXT
        )
    """)
    conn.commit()
    return conn


def fb_get(conn: sqlite3.Connection, model: str) -> Optional[Tuple[str, str, float, str]]:
    cur = conn.execute("SELECT name, description, confidence, source FROM observations WHERE model = ?", (model,))
    row = cur.fetchone()
    if not row:
        return None
    return row[0] or "", row[1] or "", float(row[2] or 0.0), row[3] or ""


def fb_put(conn: sqlite3.Connection, model: str, name: str, desc: str, conf: float, source: str) -> None:
    now = datetime.now(timezone.utc).isoformat()
    conn.execute(
        "INSERT INTO observations(model,name,description,confidence,source,updated_utc) VALUES(?,?,?,?,?,?) "
        "ON CONFLICT(model) DO UPDATE SET name=excluded.name, description=excluded.description, "
        "confidence=excluded.confidence, source=excluded.source, updated_utc=excluded.updated_utc",
        (model, name, desc, conf, source, now),
    )
    conn.commit()


def _norm_model(x: Any) -> str:
    if x is None:
        return ""
    s = str(x).strip()
    # Remove common excel weirdness
    s = s.replace("\u00a0", " ")
    s = re.sub(r"\s+", "", s)
    return s.upper()


def derive_from_series(model: str, brand: str, rules: RulePack) -> Tuple[str, str, float, str]:
    """
    Simple enrichment:
    - Try to match a known prefix from series_rules
    - Otherwise fallback to generic
    """
    m = _norm_model(model)
    if not m:
        return "", "", 0.0, "empty"

    # Find best matching prefix (longest)
    best_prefix = ""
    best_attrs: Optional[Dict[str, Any]] = None
    for pref, attrs in rules.series_map.items():
        if m.startswith(pref) and len(pref) > len(best_prefix):
            best_prefix = pref
            best_attrs = attrs

    warranty = rules.warranty_text_by_brand.get(brand.strip().lower(), "").strip()

    if best_attrs:
        equip = str(best_attrs.get("equipment_type", "")).strip()
        series_name = str(best_attrs.get("series_name", "")).strip()
        conf = float(best_attrs.get("confidence", 0.85) or 0.85)
        # Name: BRAND + (series) + model + equipment type
        pieces = [brand.upper()]
        if series_name:
            pieces.append(series_name)
        pieces.append(m)
        if equip:
            pieces.append(equip)
        name = " ".join(pieces).strip()

        # Description: keep it deterministic & readable
        desc_lines: List[str] = []
        if equip:
            desc_lines.append(f"{brand} {equip} ({m}).")
        else:
            desc_lines.append(f"{brand} HVAC equipment ({m}).")

        # Optional: add capacity hints if rules provide them
        hints = []
        for k in ("btu", "tonnage", "cabinet", "fuel", "stages"):
            v = best_attrs.get(k)
            if v:
                hints.append(f"{k}: {v}")
        if hints:
            desc_lines.append("Key attributes: " + ", ".join(hints) + ".")

        if warranty:
            desc_lines.append("Warranty: " + warranty)

        desc = "\n".join(desc_lines).strip()
        return name, desc, min(max(conf, 0.0), 1.0), "rules:series"

    # Fallback
    name = f"{brand.upper()} HVAC {m}"
    desc = f"{brand} HVAC equipment ({m})."
    if warranty:
        desc += f"\nWarranty: {warranty}"
    return name, desc, 0.60, "fallback"


def parse_args(argv: List[str]) -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--input", required=True, help="Path to skeleton XLSX")
    p.add_argument("--sheet", default=None, help="Optional sheet name")
    p.add_argument("--brand", required=True)
    p.add_argument("--manufacturer-csv", default=None)
    p.add_argument("--learned-domains", default=None)
    p.add_argument("--no-learn", action="store_true")
    p.add_argument("--model-col", default=DEFAULT_MODEL_COL)
    p.add_argument("--folder-cols", default=None)
    p.add_argument("--name-col", default=DEFAULT_NAME_COL)
    p.add_argument("--desc-col", default=DEFAULT_DESC_COL)
    p.add_argument("--bundle-mode", action="store_true")
    p.add_argument("--product-name-col", default=None)
    p.add_argument("--product-desc-col", default=None)
    p.add_argument("--part-model-col", default=None)
    p.add_argument("--start-row", type=int, default=5)
    p.add_argument("--overwrite", action="store_true")
    p.add_argument("--out", required=True, help="Output XLSX path")
    p.add_argument("--cache", default=None)

    # App-expected flags (previous v2 design)
    p.add_argument("--feedback-db", default=None)
    p.add_argument("--min-confidence", type=float, default=0.80)
    p.add_argument("--warranty-rules", default=None)
    p.add_argument("--series-rules", default=None)
    p.add_argument("--write-review", action="store_true")
    p.add_argument("--apply-overrides", action="store_true")
    p.add_argument("--save-observations", action="store_true")

    # convenience
    p.add_argument("--models", default=None, help="Optional comma-separated models to process only those")
    return p.parse_args(argv)


def main(argv: List[str]) -> int:
    args = parse_args(argv)

    in_path = Path(args.input)
    if not in_path.exists():
        print(f"[ERROR] Input not found: {in_path}", file=sys.stderr)
        return 2

    rules = load_rules(args.brand, args.warranty_rules, args.series_rules)

    fb_conn = None
    if args.feedback_db:
        fb_conn = fb_connect(args.feedback_db)

    wb = openpyxl.load_workbook(in_path)
    ws = wb[args.sheet] if args.sheet else wb.active

    model_col = (args.part_model_col or args.model_col or DEFAULT_MODEL_COL).upper()
    name_col = (args.name_col or DEFAULT_NAME_COL).upper()
    desc_col = (args.desc_col or DEFAULT_DESC_COL).upper()

    # Optional filter list
    only_models = None
    if args.models:
        only_models = set(_norm_model(x) for x in re.split(r"[,\n]+", args.models) if x.strip())

    # Ensure Review sheet
    review_ws = None
    if args.write_review:
        if "Review" in wb.sheetnames:
            review_ws = wb["Review"]
            wb.remove(review_ws)
        review_ws = wb.create_sheet("Review")
        review_ws.append(["Model", "Status", "Confidence", "Source", "NameWritten", "DescWritten"])

    processed = 0
    updated = 0
    skipped = 0

    max_row = ws.max_row or 0
    for r in range(args.start_row, max_row + 1):
        raw_model = ws[f"{model_col}{r}"].value
        model = _norm_model(raw_model)

        if not model:
            skipped += 1
            continue
        if only_models is not None and model not in only_models:
            skipped += 1
            continue

        processed += 1

        # Apply overrides first if requested
        if fb_conn and args.apply_overrides:
            ov = fb_get(fb_conn, model)
            if ov:
                ov_name, ov_desc, ov_conf, ov_source = ov
                if args.overwrite or not ws[f"{name_col}{r}"].value:
                    ws[f"{name_col}{r}"].value = ov_name
                if args.overwrite or not ws[f"{desc_col}{r}"].value:
                    ws[f"{desc_col}{r}"].value = ov_desc
                updated += 1
                if review_ws:
                    review_ws.append([model, "override", ov_conf, ov_source, bool(ov_name), bool(ov_desc)])
                continue

        # Derive
        name, desc, conf, source = derive_from_series(model, args.brand, rules)

        # If confidence below threshold and not overwriting, we still write if cells are blank
        can_write = conf >= float(args.min_confidence) or args.overwrite

        name_cell = ws[f"{name_col}{r}"]
        desc_cell = ws[f"{desc_col}{r}"]

        wrote_name = False
        wrote_desc = False

        if can_write or not name_cell.value:
            if args.overwrite or not name_cell.value:
                name_cell.value = name
                wrote_name = True

        if can_write or not desc_cell.value:
            if args.overwrite or not desc_cell.value:
                desc_cell.value = desc
                wrote_desc = True

        if wrote_name or wrote_desc:
            updated += 1

        if fb_conn and args.save_observations:
            fb_put(fb_conn, model, name, desc, conf, source)

        if review_ws:
            review_ws.append([model, "ok" if (wrote_name or wrote_desc) else "skipped", conf, source, wrote_name, wrote_desc])

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    print(f"[OK] Processed={processed} Updated={updated} Skipped={skipped} Output={out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
